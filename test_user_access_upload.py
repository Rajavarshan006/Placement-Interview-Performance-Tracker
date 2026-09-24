import openpyxl
import io
import db
from fastapi.testclient import TestClient
from app import app

client = TestClient(app)

def test_bulk_user_access_upload():
    # Create an in-memory Excel workbook with sample user emails and roles
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "User Access"
    
    # Header row
    ws.append(["gmail", "role"])
    
    # Sample rows
    ws.append(["test_student1@gmail.com", "Student"])
    ws.append(["test_recruiter1@gmail.com", "Recruiter"])
    ws.append(["test_coord1@gmail.com", "Coordinator"])
    
    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)
    
    # Send request to FastAPI endpoint
    response = client.post(
        "/api/users/upload-access",
        data={"default_role": "Student"},
        files={"file": ("test_users.xlsx", excel_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    
    assert response.status_code == 200, f"Response failed: {response.text}"
    json_data = response.json()
    assert json_data["success"] is True
    assert json_data["total_processed"] == 3
    print("API Response:", json_data)
    
    # Check DB user retrieval
    student = db.get_user_by_gmail("test_student1@gmail.com")
    assert student is not None
    assert student["role"] == "Student"
    assert student["password"] == "student123"
    
    recruiter = db.get_user_by_gmail("test_recruiter1@gmail.com")
    assert recruiter is not None
    assert recruiter["role"] == "Recruiter"
    assert recruiter["password"] == "recruiter123"
    
    coord = db.get_user_by_gmail("test_coord1@gmail.com")
    assert coord is not None
    assert coord["role"] == "Coordinator"
    assert coord["password"] == "coord123"
    
    print("All creation tests passed successfully!")

    # Pass 2: Update existing role
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append(["gmail", "role"])
    ws2.append(["test_student1@gmail.com", "Recruiter"])  # Change Student to Recruiter
    
    excel_file2 = io.BytesIO()
    wb2.save(excel_file2)
    excel_file2.seek(0)
    
    res2 = client.post(
        "/api/users/upload-access",
        data={"default_role": "Student"},
        files={"file": ("test_users2.xlsx", excel_file2, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert res2.status_code == 200
    json2 = res2.json()
    assert json2["updated_count"] == 1
    assert json2["created_count"] == 0
    
    updated_user = db.get_user_by_gmail("test_student1@gmail.com")
    assert updated_user["role"] == "Recruiter"
    print("All update tests passed successfully!")

def test_single_user_access_grant():
    # Grant single user access via endpoint
    res = client.post(
        "/api/users/grant-single-access",
        json={"gmail": "single_student@gmail.com", "role": "Student", "password": "custompwd123"}
    )
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert data["user"]["gmail"] == "single_student@gmail.com"
    assert data["user"]["role"] == "Student"
    assert data["user"]["password"] == "custompwd123"

    user_in_db = db.get_user_by_gmail("single_student@gmail.com")
    assert user_in_db is not None
    assert user_in_db["password"] == "custompwd123"
    print("Single user access test passed successfully!")

def test_bulk_upload_with_passwords():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["gmail", "role", "password"])
    ws.append(["pwd_student1@gmail.com", "Student", "MyPass123!"])
    ws.append(["pwd_student2@gmail.com", "Student", "Secure456#"])

    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    res = client.post(
        "/api/users/upload-access",
        data={"default_role": "Student"},
        files={"file": ("users_with_pwd.xlsx", excel_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True

    u1 = db.get_user_by_gmail("pwd_student1@gmail.com")
    assert u1 is not None
    assert u1["password"] == "MyPass123!"

    u2 = db.get_user_by_gmail("pwd_student2@gmail.com")
    assert u2 is not None
    assert u2["password"] == "Secure456#"
    print("Bulk upload with passwords test passed successfully!")

if __name__ == "__main__":
    test_bulk_user_access_upload()
    test_single_user_access_grant()
    test_bulk_upload_with_passwords()


