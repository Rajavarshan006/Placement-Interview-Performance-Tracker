from fastapi import FastAPI, HTTPException, status, Form, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
import os
import uvicorn
import io
import csv
import openpyxl

import db


# Initialize database on startup
db.init_db()

app = FastAPI(
    title="Placement Tracking Authentication API",
    description="Gmail & Password Authentication against SQLite 'authenticate' table"
)

# Enable CORS for React frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Allows all origins during dev
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class LoginRequest(BaseModel):
    gmail: str = Field(..., json_schema_extra={"example": "[EMAIL_ADDRESS]"})
    password: str = Field(..., json_schema_extra={"example": "admin123"})

@app.post("/api/login")
async def login(credentials: LoginRequest):
    gmail = credentials.gmail.strip()
    password = credentials.password
    
    if not gmail or not password:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Gmail and password are required"}
        )
    
    # Query database table 'authenticate' for user record
    user = db.get_user_by_gmail(gmail)
    
    # Check if user exists and comparing stored password with entered password
    if not user or user["password"] != password:
        return JSONResponse(
            status_code=status.HTTP_401_UNAUTHORIZED,
            content={
                "success": False,
                "message": "Invalid Gmail or password"
            }
        )
    
    # Passwords match -> User successfully authenticated
    return JSONResponse(
        status_code=status.HTTP_200_OK,
        content={
            "success": True,
            "message": "Logged in successfully",
            "user": {
                "uuid": user["uuid"],
                "gmail": user["gmail"],
                "role": user["role"]
            }
        }
    )

class CreateDriveRequest(BaseModel):
    company_name: str
    job_role: str
    ctc_lpa: float
    min_cgpa: float = 0.0
    allowed_branches: str = "All"
    location: str = "On Campus"
    status: str = "Active"
    deadline: str = None

@app.get("/api/users")
async def list_demo_users():
    """Helper endpoint to list available demo accounts for convenience."""
    users = db.get_all_users()
    return {"success": True, "users": users}

import io
import csv
import openpyxl

@app.get("/api/drives")
async def list_drives():
    """Endpoint to retrieve all placement drives from SQLite."""
    drives = db.get_all_drives()
    # Attach result count to each drive for convenience
    for d in drives:
        d["results_count"] = db.get_drive_results_count(d["id"])
    return {"success": True, "drives": drives}

@app.post("/api/drives")
async def create_new_drive(drive_data: CreateDriveRequest):
    """Endpoint for Coordinator to create a new placement drive."""
    if not drive_data.company_name.strip() or not drive_data.job_role.strip():
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Company name and job role are required."}
        )
    
    new_drive = db.create_drive(
        company_name=drive_data.company_name.strip(),
        job_role=drive_data.job_role.strip(),
        ctc_lpa=drive_data.ctc_lpa,
        min_cgpa=drive_data.min_cgpa,
        allowed_branches=drive_data.allowed_branches.strip(),
        location=drive_data.location.strip(),
        status=drive_data.status.strip() if drive_data.status else "Active",
        deadline=drive_data.deadline
    )
    new_drive["results_count"] = 0
    return JSONResponse(
        status_code=status.HTTP_201_CREATED,
        content={"success": True, "message": "Drive created successfully!", "drive": new_drive}
    )

from fastapi import UploadFile, File

@app.get("/api/drives/{drive_id}/results")
async def get_drive_results(drive_id: str):
    """Retrieve all student evaluation results for a specific drive."""
    results = db.get_drive_results(drive_id)
    return {"success": True, "results": results, "count": len(results)}

@app.post("/api/drives/{drive_id}/upload-results")
async def upload_drive_results(drive_id: str, file: UploadFile = File(...)):
    """
    Upload Excel (.xlsx) or CSV file containing candidate results.
    Extracts 'gmail' and 'result' columns and updates student statuses for this company drive.
    """
    filename = file.filename.lower()
    content = await file.read()
    
    rows = []
    
    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        try:
            wb = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)
            sheet = wb.active
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    rows.append([str(cell) if cell is not None else "" for cell in row])
        except Exception as e:
            return JSONResponse(
                status_code=status.HTTP_400_BAD_REQUEST,
                content={"success": False, "message": f"Error parsing Excel file: {str(e)}"}
            )
    elif filename.endswith(".csv"):
        try:
            decoded = content.decode("utf-8", errors="ignore")
            reader = csv.reader(io.StringIO(decoded))
            for r in reader:
                if any(c.strip() for c in r):
                    rows.append(r)
        except Exception as e:
            return JSONResponse(
                status_code=status.HTTP_400_BAD_REQUEST,
                content={"success": False, "message": f"Error parsing CSV file: {str(e)}"}
            )
    else:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Unsupported file format. Please upload an .xlsx or .csv file."}
        )
        
    if not rows:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Uploaded file is empty."}
        )

    # Search header row for 'gmail' and 'result' columns
    header = [str(cell).strip().lower() for cell in rows[0]]
    
    gmail_idx = -1
    result_idx = -1
    
    for idx, col in enumerate(header):
        if col in ["gmail", "email", "student email", "student gmail", "mail", "gmail_id", "email_id", "student email id"]:
            gmail_idx = idx
        elif col in ["result", "status", "round result", "round_result", "verdict", "drive status", "drive_status", "state", "selection"]:
            result_idx = idx
            
    if gmail_idx == -1:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": f"Could not find a 'gmail' or 'email' column header in the spreadsheet. Found columns: {', '.join(header)}"}
        )
        
    if result_idx == -1:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": f"Could not find a 'result' or 'status' column header in the spreadsheet. Found columns: {', '.join(header)}"}
        )

    updated_count = 0
    skipped_count = 0
    processed_records = []

    # Process data rows
    for row in rows[1:]:
        if len(row) <= max(gmail_idx, result_idx):
            skipped_count += 1
            continue
            
        gmail_val = str(row[gmail_idx]).strip()
        result_val = str(row[result_idx]).strip()
        
        # Validate gmail format basic check
        if "@" in gmail_val and result_val:
            db.upsert_student_drive_result(drive_id, gmail_val, result_val)
            updated_count += 1
            processed_records.append({"gmail": gmail_val, "result": result_val})
        else:
            skipped_count += 1

    return JSONResponse(
        status_code=status.HTTP_200_OK,
        content={
            "success": True,
            "message": f"Successfully processed {updated_count} student result records.",
            "total_rows": len(rows) - 1,
            "updated_count": updated_count,
            "skipped_count": skipped_count,
            "processed_records": processed_records
        }
    )

@app.post("/api/users/upload-access")
async def upload_user_access(
    file: UploadFile = File(...),
    default_role: str = Form("Student")
):
    """
    Upload Excel (.xlsx) or CSV file containing user email addresses.
    Grants access and creates/updates account roles in the database.
    """
    filename = file.filename.lower()
    content = await file.read()
    
    rows = []
    
    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        try:
            wb = openpyxl.load_workbook(filename=io.BytesIO(content), data_only=True)
            sheet = wb.active
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    rows.append([str(cell) if cell is not None else "" for cell in row])
        except Exception as e:
            return JSONResponse(
                status_code=status.HTTP_400_BAD_REQUEST,
                content={"success": False, "message": f"Error parsing Excel file: {str(e)}"}
            )
    elif filename.endswith(".csv"):
        try:
            decoded = content.decode("utf-8", errors="ignore")
            reader = csv.reader(io.StringIO(decoded))
            for r in reader:
                if any(c.strip() for c in r):
                    rows.append(r)
        except Exception as e:
            return JSONResponse(
                status_code=status.HTTP_400_BAD_REQUEST,
                content={"success": False, "message": f"Error parsing CSV file: {str(e)}"}
            )
    else:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Unsupported file format. Please upload an .xlsx or .csv file."}
        )
        
    if not rows:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Uploaded file is empty."}
        )

    header = [str(cell).strip().lower() for cell in rows[0]]
    
    gmail_idx = -1
    role_idx = -1
    
    for idx, col in enumerate(header):
        if col in ["gmail", "email", "student email", "user email", "mail", "gmail_id", "email_id", "student email id"]:
            gmail_idx = idx
        elif col in ["role", "user role", "access role", "account role", "type"]:
            role_idx = idx

    if gmail_idx == -1:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": f"Could not find a 'gmail' or 'email' column header in the spreadsheet. Found columns: {', '.join(header)}"}
        )

    users_to_process = []
    skipped_count = 0

    for row in rows[1:]:
        if len(row) <= gmail_idx:
            skipped_count += 1
            continue
            
        gmail_val = str(row[gmail_idx]).strip()
        role_val = str(row[role_idx]).strip() if (role_idx != -1 and len(row) > role_idx and str(row[role_idx]).strip()) else default_role
        
        if "@" in gmail_val:
            users_to_process.append({"gmail": gmail_val, "role": role_val})
        else:
            skipped_count += 1

    if not users_to_process:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "No valid Gmail addresses found in the spreadsheet."}
        )

    summary = db.bulk_grant_user_access(users_to_process)
    summary["skipped_count"] = skipped_count
    summary["total_rows"] = len(rows) - 1

    return JSONResponse(
        status_code=status.HTTP_200_OK,
        content={
            "success": True,
            "message": f"Successfully granted access to {summary['total_processed']} users ({summary['created_count']} created, {summary['updated_count']} updated).",
            **summary
        }
    )

# Serve static frontend files
PUBLIC_DIR = os.path.join(os.path.dirname(__file__), "public")
if os.path.exists(PUBLIC_DIR):
    app.mount("/static", StaticFiles(directory=PUBLIC_DIR), name="static")

@app.get("/")
async def serve_index():
    index_path = os.path.join(PUBLIC_DIR, "index.html")
    if os.path.exists(index_path):
        return FileResponse(index_path)
    return {"message": "Placement Tracking API is running."}

if __name__ == "__main__":
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=True)
